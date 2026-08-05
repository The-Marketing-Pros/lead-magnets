"""
Build smart Excel template files — white-label, sample-branded for lead magnet pack.
"""
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.formatting.rule import (
    ColorScaleRule, CellIsRule, FormulaRule, DataBarRule
)
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles.numbers import FORMAT_DATE_DATETIME
import datetime

OUT = "/Users/press/.openclaw/workspace/personas/press/lead-magnets/04-templates/files"

# ─── COLOURS ────────────────────────────────────────────────────────────────
NAVY       = "1B3A6B"
BLUE       = "0EA5E9"
GREEN      = "10B981"
GREEN_L    = "D1FAE5"
RED        = "EF4444"
RED_L      = "FEE2E2"
AMBER      = "F59E0B"
AMBER_L    = "FEF3C7"
SLATE      = "64748B"
SLATE_L    = "F1F5F9"
WHITE      = "FFFFFF"
BORDER_C   = "E2E8F0"

def side(c=BORDER_C): return Side(style="thin", color=c)
def border(all=True):
    s = side()
    if all: return Border(left=s, right=s, top=s, bottom=s)
    return Border(bottom=side())

def fill(hex): return PatternFill("solid", fgColor=hex)
def font(bold=False, color="1E293B", size=11, italic=False):
    return Font(name="Calibri", bold=bold, color=color, size=size, italic=italic)
def align(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def header_row(ws, row, cols, bg=NAVY, fg=WHITE, height=22):
    for col, text in enumerate(cols, 1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = font(bold=True, color=fg, size=10)
        c.fill = fill(bg)
        c.alignment = align("center")
        c.border = border()
    ws.row_dimensions[row].height = height

def kpi_box(ws, row, col, label, formula, color=NAVY, label_color="94A3B8"):
    lc = ws.cell(row=row, column=col, value=label)
    lc.font = Font(name="Calibri", bold=True, color=label_color, size=9)
    lc.fill = fill(color)
    lc.alignment = align("center")
    vc = ws.cell(row=row+1, column=col, value=formula)
    vc.font = Font(name="Calibri", bold=True, color=WHITE, size=18)
    vc.fill = fill(color)
    vc.alignment = align("center")
    ws.row_dimensions[row].height = 16
    ws.row_dimensions[row+1].height = 32
    return vc


# ═══════════════════════════════════════════════════════════════════════════
# 1. CREDENTIALING TRACKER
# ═══════════════════════════════════════════════════════════════════════════
def build_tracker():
    wb = openpyxl.Workbook()

    # ── Dashboard sheet ──────────────────────────────────────────────────
    dash = wb.active
    dash.title = "📊 Dashboard"
    dash.sheet_view.showGridLines = False
    dash.column_dimensions["A"].width = 2   # margin

    # Title block
    dash.merge_cells("B2:K2")
    t = dash["B2"]
    t.value = "PROVIDER CREDENTIALING TRACKER"
    t.font = Font(name="Calibri", bold=True, color=WHITE, size=20)
    t.fill = fill(NAVY)
    t.alignment = align("left", "center")
    dash.row_dimensions[2].height = 40

    dash.merge_cells("B3:K3")
    sub = dash["B3"]
    sub.value = "Riverside Medical Group  ·  Sample Data — Replace with your practice name  ·  Last updated: auto"
    sub.font = Font(name="Calibri", color="94A3B8", size=10, italic=True)
    sub.fill = fill(NAVY)
    sub.alignment = align("left", "center")
    dash.row_dimensions[3].height = 18

    # KPI row labels
    dash.row_dimensions[5].height = 16
    dash.row_dimensions[6].height = 36

    kpi_cols = {
        "B": ("TOTAL APPLICATIONS", "=COUNTA(Tracker!B5:B1000)"),
        "D": ("✅ APPROVED",          "=COUNTIF(Tracker!G5:G1000,\"Approved\")"),
        "F": ("⏳ PENDING",           "=COUNTIF(Tracker!G5:G1000,\"Pending\")"),
        "H": ("🚨 ACTION NEEDED",     "=COUNTIF(Tracker!G5:G1000,\"Info Requested\")"),
        "J": ("📆 OVERDUE FOLLOW-UPS","=COUNTIF(Tracker!I5:I1000,\"<\"&TODAY())"),
    }
    kpi_colors = {"B": NAVY, "D": "0E6B4A", "F": "1D4ED8", "H": "B91C1C", "J": "92400E"}
    for col, (label, formula) in kpi_cols.items():
        lc = dash[f"{col}5"]
        lc.value = label
        lc.font = Font(name="Calibri", bold=True, color="BFDBFE", size=8)
        lc.fill = fill(kpi_colors[col])
        lc.alignment = align("center")
        vc = dash[f"{col}6"]
        vc.value = formula
        vc.font = Font(name="Calibri", bold=True, color=WHITE, size=22)
        vc.fill = fill(kpi_colors[col])
        vc.alignment = align("center")
        # span 2 cols each
        next_col = chr(ord(col)+1)
        dash.merge_cells(f"{col}5:{next_col}5")
        dash.merge_cells(f"{col}6:{next_col}6")

    # Instructions box
    dash.merge_cells("B8:K8")
    dash["B8"].value = "→  Go to the Tracker tab to enter applications.  Status dropdowns, overdue flags, and all KPIs above update automatically."
    dash["B8"].font = Font(name="Calibri", color="1E40AF", size=10, italic=True)
    dash["B8"].fill = fill("EFF6FF")
    dash["B8"].alignment = align("left", "center")
    dash.row_dimensions[8].height = 20

    for col in "BCDEFGHIJK":
        dash.column_dimensions[col].width = 14

    # ── Tracker sheet ────────────────────────────────────────────────────
    tr = wb.create_sheet("Tracker")
    tr.sheet_view.showGridLines = False
    tr.freeze_panes = "A5"

    # Title
    tr.merge_cells("A1:N1")
    t = tr["A1"]
    t.value = "CREDENTIALING APPLICATION TRACKER"
    t.font = Font(name="Calibri", bold=True, color=WHITE, size=14)
    t.fill = fill(NAVY)
    t.alignment = align("left", "center")
    tr.row_dimensions[1].height = 32

    tr.merge_cells("A2:N2")
    s = tr["A2"]
    s.value = "Enter one row per payer per provider. Status dropdown auto-updates KPIs on Dashboard. Overdue = Follow-Up Date < Today."
    s.font = Font(name="Calibri", color=SLATE, size=9, italic=True)
    s.fill = fill(SLATE_L)
    s.alignment = align("left", "center")
    tr.row_dimensions[2].height = 16

    tr.row_dimensions[3].height = 8  # spacer

    headers = [
        "PROVIDER NAME", "NPI", "PAYER", "APP TYPE",
        "DATE SUBMITTED", "DAYS SINCE SUB", "STATUS", "FOLLOW-UP DATE",
        "OVERDUE?", "TARGET EFFECTIVE", "ACTUAL EFFECTIVE",
        "CONTRACT RATE", "ANNUAL IMPACT", "NOTES"
    ]
    header_row(tr, 4, headers, height=24)

    col_widths = [22, 13, 22, 16, 15, 15, 16, 15, 12, 16, 16, 14, 14, 28]
    for i, w in enumerate(col_widths, 1):
        tr.column_dimensions[get_column_letter(i)].width = w

    # Sample data rows + formulas
    samples = [
        ["Dr. Sarah Chen, MD", "1234567890", "Medicare (PECOS)", "Initial Enrollment",
         datetime.date(2026, 6, 3), None, "Pending", datetime.date(2026, 6, 24),
         None, None, None, None, None, "PECOS #45921"],
        ["Dr. Sarah Chen, MD", "1234567890", "BCBS Tennessee", "Commercial",
         datetime.date(2026, 6, 5), None, "Info Requested", datetime.date(2026, 6, 7),
         None, None, None, None, None, "AIR: needs updated CV"],
        ["Dr. Sarah Chen, MD", "1234567890", "Aetna", "Commercial",
         datetime.date(2026, 5, 28), None, "Approved", None,
         None, datetime.date(2026, 7, 1), datetime.date(2026, 7, 1), "105% Medicare", 840000, "Test claim OK"],
        ["Dr. Marcus Webb, DO", "9876543210", "UnitedHealthcare", "Commercial",
         datetime.date(2026, 6, 10), None, "Pending", datetime.date(2026, 7, 1),
         None, None, None, None, None, "Portal ref UHC-88241"],
        ["Dr. Marcus Webb, DO", "9876543210", "Cigna", "Commercial",
         datetime.date(2026, 6, 10), None, "Submitted", datetime.date(2026, 7, 1),
         None, None, None, None, None, "Fax confirmation received"],
    ]

    for i, row_data in enumerate(samples, 5):
        row_num = i
        for j, val in enumerate(row_data, 1):
            c = tr.cell(row=row_num, column=j, value=val)
            c.font = font(size=10)
            c.alignment = align(wrap=True)
            c.border = border()
            if j in (5, 8, 10, 11) and isinstance(val, datetime.date):
                c.number_format = "MM/DD/YYYY"
        # Days since submitted formula (col F = 6)
        days_c = tr.cell(row=row_num, column=6)
        days_c.value = f'=IF(E{row_num}="","",TODAY()-E{row_num})'
        days_c.font = Font(name="Calibri", bold=True, color=NAVY, size=10)
        days_c.number_format = '0" days"'
        days_c.alignment = align("center")
        # Overdue formula (col I = 9)
        ov_c = tr.cell(row=row_num, column=9)
        ov_c.value = f'=IF(H{row_num}="","",IF(AND(H{row_num}<TODAY(),G{row_num}<>"Approved"),"⚠️ OVERDUE","OK"))'
        ov_c.font = Font(name="Calibri", bold=True, size=10)
        ov_c.alignment = align("center")
        tr.row_dimensions[row_num].height = 20

    # Add 45 blank data rows with formulas
    for i in range(len(samples)+5, 55):
        days_c = tr.cell(row=i, column=6)
        days_c.value = f'=IF(E{i}="","",TODAY()-E{i})'
        days_c.number_format = '0" days"'
        days_c.font = Font(name="Calibri", bold=True, color=NAVY, size=10)
        days_c.alignment = align("center")
        ov_c = tr.cell(row=i, column=9)
        ov_c.value = f'=IF(H{i}="","",IF(AND(H{i}<TODAY(),G{i}<>"Approved"),"⚠️ OVERDUE","OK"))'
        ov_c.font = Font(name="Calibri", bold=True, size=10)
        ov_c.alignment = align("center")
        for j in range(1, 15):
            c = tr.cell(row=i, column=j)
            if not c.value:
                c.border = border()
            c.font = font(size=10)
            c.alignment = align(wrap=True)
        tr.row_dimensions[i].height = 18

    # Status dropdown validation
    dv_status = DataValidation(
        type="list",
        formula1='"Not Started,Submitted,Pending,Info Requested,Approved,Effective,Denied,Re-cred Due"',
        allow_blank=True,
        showDropDown=False
    )
    dv_status.prompt = "Select application status"
    dv_status.promptTitle = "Status"
    tr.add_data_validation(dv_status)
    dv_status.sqref = "G5:G1000"

    # App type dropdown
    dv_type = DataValidation(
        type="list",
        formula1='"Initial Enrollment,Re-credentialing,Commercial,Medicare,Medicaid,Medicare Advantage,Workers Comp,Tricare"',
        allow_blank=True
    )
    tr.add_data_validation(dv_type)
    dv_type.sqref = "D5:D1000"

    # ── Conditional formatting ──────────────────────────────────────────
    # Status = Approved → green
    tr.conditional_formatting.add("G5:G1000",
        FormulaRule(formula=['G5="Approved"'], fill=fill(GREEN_L),
                    font=Font(name="Calibri", bold=True, color="065F46", size=10)))
    # Status = Info Requested → red
    tr.conditional_formatting.add("G5:G1000",
        FormulaRule(formula=['G5="Info Requested"'], fill=fill(RED_L),
                    font=Font(name="Calibri", bold=True, color="991B1B", size=10)))
    # Status = Pending → amber
    tr.conditional_formatting.add("G5:G1000",
        FormulaRule(formula=['G5="Pending"'], fill=fill(AMBER_L),
                    font=Font(name="Calibri", bold=True, color="92400E", size=10)))
    # Overdue flag → red fill on whole row
    tr.conditional_formatting.add("A5:N1000",
        FormulaRule(formula=['$I5="⚠️ OVERDUE"'], fill=fill("FFF1F2"),
                    font=Font(name="Calibri", color="991B1B", size=10)))
    # Overdue cell itself
    tr.conditional_formatting.add("I5:I1000",
        FormulaRule(formula=['I5="⚠️ OVERDUE"'], fill=fill(RED_L),
                    font=Font(name="Calibri", bold=True, color="991B1B", size=10)))
    # Days since > 90 → amber warning on days col
    tr.conditional_formatting.add("F5:F1000",
        CellIsRule(operator="greaterThan", formula=["90"], fill=fill(AMBER_L),
                   font=Font(name="Calibri", bold=True, color="92400E", size=10)))

    # ── Status legend ────────────────────────────────────────────────────
    leg = wb.create_sheet("📖 Legend & Help")
    leg.sheet_view.showGridLines = False
    leg.column_dimensions["A"].width = 2
    leg.column_dimensions["B"].width = 24
    leg.column_dimensions["C"].width = 50

    leg.merge_cells("B2:C2")
    leg["B2"].value = "STATUS COLOUR LEGEND"
    leg["B2"].font = font(bold=True, color=WHITE, size=13)
    leg["B2"].fill = fill(NAVY)
    leg["B2"].alignment = align("center")
    leg.row_dimensions[2].height = 28

    legends = [
        ("Not Started",    "94A3B8", "F8FAFC", "Application not yet submitted"),
        ("Submitted",      "1E40AF", "EFF6FF", "Application sent — awaiting acknowledgment"),
        ("Pending",        "92400E", AMBER_L,  "In payer review queue — follow up every 3 weeks"),
        ("Info Requested", "991B1B", RED_L,    "⚠️ Payer needs documents — respond within 48 hours"),
        ("Approved",       "065F46", GREEN_L,  "✅ Credentialing approved — confirm effective date"),
        ("Effective",      "065F46", "D1FAE5", "Live and billing — verify test claim processed"),
        ("Denied",         "7F1D1D", "FEE2E2", "Rejected — see notes for reason and appeal plan"),
        ("Re-cred Due",    "92400E", AMBER_L,  "Re-credentialing deadline within 90 days"),
    ]
    for i, (status, fc, bg, desc) in enumerate(legends, 4):
        sc = leg.cell(row=i, column=2, value=status)
        sc.font = Font(name="Calibri", bold=True, color=fc, size=10)
        sc.fill = fill(bg)
        sc.alignment = align("center")
        sc.border = border()
        dc = leg.cell(row=i, column=3, value=desc)
        dc.font = font(size=10)
        dc.fill = fill(bg)
        dc.alignment = align(wrap=True)
        dc.border = border()
        leg.row_dimensions[i].height = 18

    wb.save(f"{OUT}/01-credentialing-tracker.xlsx")
    print("✓ 01-credentialing-tracker.xlsx")


# ═══════════════════════════════════════════════════════════════════════════
# 2. PAYER PRIORITY MATRIX
# ═══════════════════════════════════════════════════════════════════════════
def build_payer_matrix():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "🎯 Priority Matrix"
    ws.sheet_view.showGridLines = False

    # Title
    ws.merge_cells("A1:L1")
    t = ws["A1"]
    t.value = "PAYER PRIORITY MATRIX  ·  Riverside Medical Group  (Sample)"
    t.font = Font(name="Calibri", bold=True, color=WHITE, size=16)
    t.fill = fill(NAVY)
    t.alignment = align("left", "center")
    ws.row_dimensions[1].height = 38

    ws.merge_cells("A2:L2")
    s = ws["A2"]
    s.value = "Score each payer 1–5 per factor. Total and Rank update automatically. Higher score = apply first."
    s.font = Font(name="Calibri", color="1D4ED8", size=10, italic=True)
    s.fill = fill("EFF6FF")
    s.alignment = align("left", "center")
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 8  # spacer

    headers = [
        "PAYER", "TYPE",
        "PATIENT VOLUME\n(1–5)",
        "REIMBURSEMENT\nRATE (1–5)",
        "APPLICATION\nEASE (1–5)",
        "TIMELINE\nSPEED (1–5)",
        "IN-NETWORK\nDEMAND (1–5)",
        "STRATEGIC\nIMPORTANCE (1–5)",
        "TOTAL\nSCORE",
        "PRIORITY\nRANK",
        "SCORE\nVISUAL",
        "NOTES / STRATEGY"
    ]
    header_row(ws, 4, headers, height=36)

    col_widths = [22, 14, 14, 14, 14, 14, 14, 16, 10, 10, 16, 32]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    payers = [
        ("Medicare",            "Government", 5, 3, 4, 2, 5, 5, "CMS PECOS enrollment — required for most referrals"),
        ("Medicaid — [State]",  "Government", 4, 2, 3, 2, 4, 4, "State portal — check state-specific requirements"),
        ("BCBS [Region]",       "Commercial", 4, 4, 3, 3, 4, 4, "Largest commercial network; negotiate rate upfront"),
        ("Aetna",               "Commercial", 3, 4, 3, 3, 3, 4, "Strong employer base; 90–120 day timeline in hot markets"),
        ("UnitedHealthcare",    "Commercial", 4, 4, 2, 2, 4, 5, "Largest US commercial payer — start early (90–120d)"),
        ("Cigna",               "Commercial", 3, 4, 3, 3, 3, 3, "Solid employer coverage; negotiate before signing"),
        ("Humana",              "Commercial", 3, 3, 3, 3, 2, 3, "Stronger in Medicare Advantage markets"),
        ("Tricare",             "Government", 2, 3, 3, 3, 2, 3, "Military — evaluate if near base"),
        ("[Add Payer]",         "",           "", "", "", "", "", "", ""),
        ("[Add Payer]",         "",           "", "", "", "", "", "", ""),
        ("[Add Payer]",         "",           "", "", "", "", "", "", ""),
    ]

    for i, row_data in enumerate(payers, 5):
        row_num = i
        name, ptype = row_data[0], row_data[1]
        scores = row_data[2:8]
        note = row_data[8]

        # Payer name
        nc = ws.cell(row=row_num, column=1, value=name)
        nc.font = Font(name="Calibri", bold=True, color=NAVY, size=10)
        nc.border = border()
        nc.alignment = align()

        # Type
        tc = ws.cell(row=row_num, column=2, value=ptype)
        tc.font = font(size=10)
        tc.border = border()
        tc.alignment = align("center")

        # Score columns C–H (3–8)
        for j, score in enumerate(scores, 3):
            c = ws.cell(row=row_num, column=j, value=score if score != "" else None)
            c.font = Font(name="Calibri", bold=True, color=NAVY, size=11)
            c.border = border()
            c.alignment = align("center")

        # Total formula (col I = 9) sum of C:H
        col_letters = [get_column_letter(j) for j in range(3, 9)]
        total_formula = f'=SUM({col_letters[0]}{row_num}:{col_letters[-1]}{row_num})'
        tot = ws.cell(row=row_num, column=9, value=total_formula)
        tot.font = Font(name="Calibri", bold=True, color=WHITE, size=12)
        tot.fill = fill(NAVY)
        tot.border = border()
        tot.alignment = align("center")

        # Rank formula (col J = 10)
        rank = ws.cell(row=row_num, column=10)
        rank.value = f'=IF(I{row_num}=0,"—",RANK(I{row_num},$I$5:$I$15,0))'
        rank.font = Font(name="Calibri", bold=True, color=NAVY, size=11)
        rank.border = border()
        rank.alignment = align("center")

        # Score visual bar (col K = 11) using REPT
        bar = ws.cell(row=row_num, column=11)
        bar.value = f'=IF(I{row_num}=0,"",REPT("█",ROUND(I{row_num}/2,0))&" "&I{row_num})'
        bar.font = Font(name="Calibri", color=BLUE, size=9)
        bar.border = border()
        bar.alignment = align()

        # Notes (col L = 12)
        nc2 = ws.cell(row=row_num, column=12, value=note)
        nc2.font = font(size=10, italic=True, color=SLATE)
        nc2.border = border()
        nc2.alignment = align(wrap=True)

        ws.row_dimensions[row_num].height = 20

    # Dropdown for payer type
    dv = DataValidation(type="list",
        formula1='"Government,Commercial,Medicare Advantage,Workers Comp,Tricare,Other"',
        allow_blank=True)
    ws.add_data_validation(dv)
    dv.sqref = "B5:B1000"

    # Score dropdowns 1–5
    dv_score = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
    ws.add_data_validation(dv_score)
    dv_score.sqref = "C5:H1000"

    # Conditional formatting: total score colour scale
    ws.conditional_formatting.add("I5:I15",
        ColorScaleRule(start_type="min", start_color="FEE2E2",
                       mid_type="percentile", mid_value=50, mid_color="FEF3C7",
                       end_type="max", end_color="D1FAE5"))

    # Top rank (rank = 1) → gold highlight on row
    ws.conditional_formatting.add("A5:L15",
        FormulaRule(formula=["$J5=1"], fill=fill("FEF9C3"),
                    font=Font(name="Calibri", bold=True, color="713F12", size=10)))

    # Score cells: 5 = green, 1-2 = red
    for col in "CDEFGH":
        ws.conditional_formatting.add(f"{col}5:{col}15",
            CellIsRule(operator="equal", formula=["5"], fill=fill(GREEN_L),
                       font=Font(name="Calibri", bold=True, color="065F46", size=11)))
        ws.conditional_formatting.add(f"{col}5:{col}15",
            CellIsRule(operator="lessThanOrEqual", formula=["2"], fill=fill(RED_L),
                       font=Font(name="Calibri", bold=True, color="991B1B", size=11)))

    # ── Scoring guide sheet ──────────────────────────────────────────────
    guide = wb.create_sheet("📖 Scoring Guide")
    guide.sheet_view.showGridLines = False
    guide.column_dimensions["A"].width = 2
    guide.column_dimensions["B"].width = 20
    guide.column_dimensions["C"].width = 26
    guide.column_dimensions["D"].width = 26
    guide.column_dimensions["E"].width = 26

    guide.merge_cells("B2:E2")
    guide["B2"].value = "SCORING GUIDE — HOW TO RATE EACH FACTOR"
    guide["B2"].font = font(bold=True, color=WHITE, size=13)
    guide["B2"].fill = fill(NAVY)
    guide["B2"].alignment = align("center")
    guide.row_dimensions[2].height = 28

    factors = [
        ("Patient Volume", "1 = <5% of expected pts", "3 = 15-25% of pts", "5 = >30% of expected pts"),
        ("Reimbursement Rate", "1 = <80% Medicare", "3 = 90-110% Medicare", "5 = >130% Medicare"),
        ("Application Ease", "1 = Complex/error-prone", "3 = Standard process", "5 = Simple/online portal"),
        ("Timeline Speed", "1 = 120+ days typical", "3 = 60-90 days typical", "5 = <45 days typical"),
        ("In-Network Demand", "1 = Rarely requested", "3 = Moderate patient ask", "5 = Frequently requested"),
        ("Strategic Importance", "1 = Nice to have", "3 = Important for growth", "5 = Critical — must have"),
    ]
    header_row(guide, 4, ["FACTOR", "SCORE 1–2 (Low)", "SCORE 3 (Medium)", "SCORE 4–5 (High)"], bg=SLATE)
    for i, (factor, low, mid, high) in enumerate(factors, 5):
        cells = [factor, low, mid, high]
        fills = [SLATE_L, RED_L, AMBER_L, GREEN_L]
        for j, (val, bg) in enumerate(zip(cells, fills), 2):
            c = guide.cell(row=i, column=j, value=val)
            c.font = font(size=10, bold=(j==2))
            c.fill = fill(bg)
            c.border = border()
            c.alignment = align(wrap=True)
        guide.row_dimensions[i].height = 22

    wb.save(f"{OUT}/02-payer-priority-matrix.xlsx")
    print("✓ 02-payer-priority-matrix.xlsx")


# ═══════════════════════════════════════════════════════════════════════════
# 3. CREDENTIALING TIMELINE
# ═══════════════════════════════════════════════════════════════════════════
def build_timeline():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "📅 Timeline"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"

    # Title
    ws.merge_cells("A1:J1")
    t = ws["A1"]
    t.value = "PROVIDER CREDENTIALING TIMELINE  ·  Riverside Medical Group  (Sample)"
    t.font = Font(name="Calibri", bold=True, color=WHITE, size=15)
    t.fill = fill(NAVY)
    t.alignment = align("left", "center")
    ws.row_dimensions[1].height = 36

    # Start date input
    ws.merge_cells("A2:C2")
    ws["A2"].value = "📅  START DATE  →"
    ws["A2"].font = Font(name="Calibri", bold=True, color=WHITE, size=11)
    ws["A2"].fill = fill(BLUE)
    ws["A2"].alignment = align("right", "center")

    ws["D2"].value = datetime.date.today()
    ws["D2"].font = Font(name="Calibri", bold=True, color=NAVY, size=12)
    ws["D2"].fill = fill("FEF9C3")
    ws["D2"].number_format = "MM/DD/YYYY"
    ws["D2"].border = Border(bottom=Side(style="medium", color=AMBER), left=Side(style="medium", color=AMBER), right=Side(style="medium", color=AMBER), top=Side(style="medium", color=AMBER))
    ws["D2"].alignment = align("center")

    ws.merge_cells("E2:F2")
    ws["E2"].value = "← Change this date. All milestone dates update automatically."
    ws["E2"].font = Font(name="Calibri", color=AMBER, size=10, italic=True)
    ws["E2"].alignment = align("left", "center")
    ws.row_dimensions[2].height = 28

    # Progress summary
    ws.merge_cells("A3:C3")
    ws["A3"].value = "OVERALL PROGRESS"
    ws["A3"].font = Font(name="Calibri", bold=True, color=WHITE, size=10)
    ws["A3"].fill = fill(SLATE)
    ws["A3"].alignment = align("center")

    ws["D3"].value = '=TEXT(COUNTIF(H6:H100,"Complete")/COUNTA(A6:A100),"0%")&" complete  ("&COUNTIF(H6:H100,"Complete")&" of "&COUNTA(A6:A100)&" tasks)"'
    ws["D3"].font = Font(name="Calibri", bold=True, color=NAVY, size=11)
    ws["D3"].fill = fill(GREEN_L)
    ws["D3"].alignment = align("left", "center")
    ws.merge_cells("D3:J3")
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 8  # spacer

    headers = ["PHASE", "TASK", "OWNER", "WEEK TARGET",
               "TARGET DATE", "TODAY?", "DAYS TO DUE", "STATUS", "COMPLETE DATE", "NOTES"]
    header_row(ws, 5, headers, height=26)

    col_widths = [20, 40, 18, 12, 14, 10, 12, 16, 14, 28]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Timeline data
    phases = [
        # (phase, task, owner, week_offset_days)
        ("Pre-App Prep", "Collect all provider documents",                "Practice Mgr",   7),
        ("Pre-App Prep", "Verify NPI — check NPPES registry",             "Practice Mgr",   7),
        ("Pre-App Prep", "Create / update CAQH ProView profile",          "Cred Specialist",7),
        ("Pre-App Prep", "Upload all documents to CAQH",                  "Cred Specialist",9),
        ("Pre-App Prep", "Grant payer authorizations in CAQH",            "Cred Specialist",9),
        ("Pre-App Prep", "Complete CAQH attestation",                     "Provider",       10),
        ("Pre-App Prep", "Build payer priority matrix — rank applications","Practice Mgr",   7),
        ("Pre-App Prep", "Confirm malpractice coverage is adequate",       "Practice Mgr",   14),
        ("Submission",   "Submit Medicare PECOS enrollment",               "Cred Specialist",14),
        ("Submission",   "Submit Medicaid enrollment (state portal)",       "Cred Specialist",14),
        ("Submission",   "Submit BCBS application",                        "Cred Specialist",14),
        ("Submission",   "Submit Aetna application",                       "Cred Specialist",21),
        ("Submission",   "Submit UnitedHealthcare application",             "Cred Specialist",21),
        ("Submission",   "Submit Cigna application",                        "Cred Specialist",21),
        ("Submission",   "Submit additional payer applications",            "Cred Specialist",28),
        ("Submission",   "Log all confirmation numbers + contacts",         "Cred Specialist",28),
        ("Follow-Up",    "Follow up with all payers — 1st check",          "Cred Specialist",35),
        ("Follow-Up",    "Re-attest CAQH if approaching 120-day expiry",   "Provider",       56),
        ("Follow-Up",    "Follow up with all payers — 2nd check",          "Cred Specialist",56),
        ("Follow-Up",    "Follow up with all payers — 3rd check",          "Cred Specialist",77),
        ("Follow-Up",    "Escalate any application >90 days no movement",  "Cred Specialist",91),
        ("Follow-Up",    "Follow up with all payers — 4th check",          "Cred Specialist",98),
        ("Go-Live",      "Get written effective date — each payer",        "Cred Specialist",0),
        ("Go-Live",      "Add provider to PM / billing system",            "Billing Mgr",    0),
        ("Go-Live",      "Verify EDI/ERA enrollment active",               "Billing Mgr",    3),
        ("Go-Live",      "Submit test claim — verify it processes",        "Billing Mgr",    3),
        ("Go-Live",      "Confirm fee schedule rates match contract",       "Billing Mgr",    14),
        ("Go-Live",      "Update payer provider directory listings",        "Cred Specialist",30),
        ("Maintenance",  "Set re-credentialing reminders (2–3 yr cycles)", "Cred Specialist",0),
        ("Maintenance",  "Set CAQH re-attestation recurring (90 days)",    "Provider",       0),
        ("Maintenance",  "Set license + DEA renewal reminders (60d out)",  "Practice Mgr",   0),
        ("Maintenance",  "Quarterly underpayment audit vs. contract",       "Billing Mgr",    90),
        ("Maintenance",  "Annual payer contract review",                    "Practice Mgr",   365),
    ]

    phase_colors = {
        "Pre-App Prep": ("1D4ED8", "DBEAFE"),
        "Submission":   ("92400E", "FEF3C7"),
        "Follow-Up":    ("065F46", "D1FAE5"),
        "Go-Live":      ("6B21A8", "F3E8FF"),
        "Maintenance":  ("BE123C", "FFE4E6"),
    }

    for i, (phase, task, owner, offset) in enumerate(phases, 6):
        row_num = i
        fc, bg = phase_colors.get(phase, (SLATE, SLATE_L))

        pc = ws.cell(row=row_num, column=1, value=phase)
        pc.font = Font(name="Calibri", bold=True, color=fc, size=10)
        pc.fill = fill(bg)
        pc.border = border()
        pc.alignment = align("center")

        tc = ws.cell(row=row_num, column=2, value=task)
        tc.font = font(size=10)
        tc.border = border()
        tc.alignment = align(wrap=True)

        oc = ws.cell(row=row_num, column=3, value=owner)
        oc.font = font(size=10, italic=True, color=SLATE)
        oc.border = border()
        oc.alignment = align("center")

        wk = max(1, round(offset/7)) if offset > 0 else "On Approval"
        wkc = ws.cell(row=row_num, column=4, value=f"Week {wk}" if isinstance(wk, int) else wk)
        wkc.font = font(size=10)
        wkc.border = border()
        wkc.alignment = align("center")

        # Target date formula off $D$2
        datec = ws.cell(row=row_num, column=5)
        if offset > 0:
            datec.value = f"=$D$2+{offset}"
            datec.number_format = "MM/DD/YYYY"
        else:
            datec.value = "On Approval"
        datec.font = font(size=10)
        datec.border = border()
        datec.alignment = align("center")

        # TODAY indicator
        today_c = ws.cell(row=row_num, column=6)
        today_c.value = f'=IF(E{row_num}="On Approval","",IF(AND(TODAY()>=E{row_num},TODAY()<E{row_num}+7),"▶ NOW",""))'
        today_c.font = Font(name="Calibri", bold=True, color=AMBER, size=11)
        today_c.border = border()
        today_c.alignment = align("center")

        # Days to due
        days_c = ws.cell(row=row_num, column=7)
        days_c.value = f'=IF(E{row_num}="On Approval","",E{row_num}-TODAY())'
        days_c.number_format = '0" d"'
        days_c.font = Font(name="Calibri", bold=True, size=10)
        days_c.border = border()
        days_c.alignment = align("center")

        # Status (col H = 8) — dropdown
        sc = ws.cell(row=row_num, column=8, value="Not Started")
        sc.font = font(size=10)
        sc.border = border()
        sc.alignment = align("center")

        # Complete date (col I = 9)
        cc = ws.cell(row=row_num, column=9)
        cc.font = font(size=10)
        cc.number_format = "MM/DD/YYYY"
        cc.border = border()
        cc.alignment = align("center")

        # Notes (col J = 10)
        nc = ws.cell(row=row_num, column=10)
        nc.font = font(size=10, italic=True, color=SLATE)
        nc.border = border()
        nc.alignment = align(wrap=True)

        ws.row_dimensions[row_num].height = 20

    last_row = len(phases) + 5

    # Status dropdown
    dv = DataValidation(type="list",
        formula1='"Not Started,In Progress,Complete,Blocked,Skipped,Recurring"',
        allow_blank=True)
    ws.add_data_validation(dv)
    dv.sqref = f"H6:H{last_row}"

    # Conditional formatting
    # Complete → green row
    ws.conditional_formatting.add(f"A6:J{last_row}",
        FormulaRule(formula=[f"$H6=\"Complete\""], fill=fill(GREEN_L),
                    font=Font(name="Calibri", color="065F46", size=10, strike=True)))
    # In Progress → blue
    ws.conditional_formatting.add(f"A6:J{last_row}",
        FormulaRule(formula=[f"$H6=\"In Progress\""], fill=fill("EFF6FF"),
                    font=Font(name="Calibri", bold=True, color="1D4ED8", size=10)))
    # Blocked → red
    ws.conditional_formatting.add(f"A6:J{last_row}",
        FormulaRule(formula=[f"$H6=\"Blocked\""], fill=fill(RED_L),
                    font=Font(name="Calibri", bold=True, color="991B1B", size=10)))
    # Today indicator row highlight
    ws.conditional_formatting.add(f"A6:J{last_row}",
        FormulaRule(formula=[f"$F6=\"▶ NOW\""], fill=fill("FFFBEB"),
                    font=Font(name="Calibri", bold=True, color="78350F", size=10)))
    # Days to due < 0 (overdue) → red on days col
    ws.conditional_formatting.add(f"G6:G{last_row}",
        CellIsRule(operator="lessThan", formula=["0"], fill=fill(RED_L),
                   font=Font(name="Calibri", bold=True, color="991B1B", size=10)))
    # Days to due 0–7 → amber
    ws.conditional_formatting.add(f"G6:G{last_row}",
        CellIsRule(operator="between", formula=["0","7"], fill=fill(AMBER_L),
                   font=Font(name="Calibri", bold=True, color="92400E", size=10)))

    wb.save(f"{OUT}/05-credentialing-timeline.xlsx")
    print("✓ 05-credentialing-timeline.xlsx")


if __name__ == "__main__":
    build_tracker()
    build_payer_matrix()
    build_timeline()
    print("\nAll done.")
